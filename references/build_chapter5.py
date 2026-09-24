from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "references" / "CL-AIGC智能创作与管理平台V1.0设计说明书-第5章.docx"
LOGIN_IMAGE = ROOT / "references" / "chapter5-screenshots" / "01-login.png"
WORKSPACE_IMAGE = ROOT / "references" / "chapter5-screenshots" / "04-workspace-home-owner.png"

# standard_business_brief preset, with a named Chinese-font override.
PAGE_WIDTH = Inches(8.5)
PAGE_HEIGHT = Inches(11)
MARGIN = Inches(1)
CONTENT_DXA = 9360
TABLE_INDENT_DXA = 120
CELL_MARGIN = {"top": 80, "bottom": 80, "start": 120, "end": 120}

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "172B4D"
GRAY = "5E6C84"
LIGHT_GRAY = "F2F4F7"
PALE_BLUE = "E8EEF5"
PALE_GOLD = "FFF4CE"
GREEN = "2F6B4F"
RED = "9B1C1C"
WHITE = "FFFFFF"
DOC_FONT = "Arial Unicode MS"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, margins=CELL_MARGIN):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in margins.items():
        tag = "start" if side == "start" else "end" if side == "end" else side
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths[idx])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_run_font(run, size=11, bold=None, color=None, italic=None):
    run.font.name = DOC_FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), DOC_FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), DOC_FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), DOC_FONT)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def style_paragraph_runs(paragraph, size=11, bold=None, color=None):
    for run in paragraph.runs:
        set_run_font(run, size=size, bold=bold, color=color)


def add_body(doc, text, bold_prefix=None, color=None, keep_with_next=False):
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.keep_with_next = keep_with_next
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_run_font(r1, bold=True, color=color or INK)
        r2 = p.add_run(text[len(bold_prefix):])
        set_run_font(r2, color=color)
    else:
        r = p.add_run(text)
        set_run_font(r, color=color)
    return p


def add_step(doc, number, title, detail):
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.first_line_indent = Inches(-0.18)
    p.paragraph_format.keep_together = True
    tag = p.add_run(f"步骤{number}｜{title}  ")
    set_run_font(tag, bold=True, color=DARK_BLUE)
    detail_run = p.add_run(detail)
    set_run_font(detail_run)
    return p


def add_note(doc, label, text, kind="note"):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    set_table_geometry(table, [CONTENT_DXA])
    cell = table.cell(0, 0)
    set_cell_shading(cell, PALE_GOLD if kind == "pending" else PALE_BLUE)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    r1 = p.add_run(f"{label}：")
    set_run_font(r1, bold=True, color=RED if kind == "pending" else DARK_BLUE)
    r2 = p.add_run(text)
    set_run_font(r2)
    return table


def add_table(doc, headers, rows, widths, header_fill=LIGHT_GRAY):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    hdr = table.rows[0]
    for idx, header in enumerate(headers):
        cell = hdr.cells[idx]
        set_cell_shading(cell, header_fill)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(header)
        set_run_font(r, size=9.5, bold=True, color=INK)
    for row_data in rows:
        row = table.add_row()
        for idx, value in enumerate(row_data):
            cell = row.cells[idx]
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(str(value))
            set_run_font(r, size=9.5)
            if idx == 0 and len(headers) == 2:
                r.bold = True
                r.font.color.rgb = RGBColor.from_string(DARK_BLUE)
    for row in table.rows:
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
    return table


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(text)
    set_run_font(r, size=9, color=GRAY)
    return p


def add_picture(doc, path, width=6.2):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(str(path), width=Inches(width))
    return p


def configure_styles(doc):
    section = doc.sections[0]
    section.page_width = PAGE_WIDTH
    section.page_height = PAGE_HEIGHT
    section.top_margin = MARGIN
    section.bottom_margin = MARGIN
    section.left_margin = MARGIN
    section.right_margin = MARGIN
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = DOC_FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), DOC_FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), DOC_FONT)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), DOC_FONT)
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    heading_specs = {
        "Heading 1": (16, BLUE, 16, 8),
        "Heading 2": (13, BLUE, 12, 6),
        "Heading 3": (12, DARK_BLUE, 8, 4),
    }
    for name, (size, color, before, after) in heading_specs.items():
        style = styles[name]
        style.font.name = DOC_FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), DOC_FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), DOC_FONT)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), DOC_FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    # Running header and footer, memo_masthead pattern without a decorative rule.
    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hr = hp.add_run("CL-AIGC智能创作与管理平台V1.0设计说明书")
    set_run_font(hr, size=8.5, color=GRAY)
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fr = fp.add_run("第5章 账号与访问管理  ·  ")
    set_run_font(fr, size=8.5, color=GRAY)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run = fp.add_run()
    run._r.extend([fld_begin, instr, fld_end])


def build_document():
    doc = Document()
    configure_styles(doc)

    # memo_masthead opening block
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run("设计说明书 · 账号与权限")
    set_run_font(r, size=10, bold=True, color=GRAY)

    title = doc.add_paragraph()
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(4)
    tr = title.add_run("第5章  账号与访问管理")
    set_run_font(tr, size=23, bold=True, color=INK)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(14)
    sr = subtitle.add_run("CL-AIGC智能创作与管理平台V1.0")
    set_run_font(sr, size=13, color=GRAY)

    add_table(
        doc,
        ["项目", "说明"],
        [
            ["核实范围", "用户登录、企业主账号、内部子账号、Harson-Base用户映射、工作区访问和退出登录"],
            ["核实方式", "当前项目源代码核对 + 本地浏览器页面核验"],
            ["数据边界", "未读取或展示secure CSV内容、密码、Token、加密凭据；未执行账号创建、绑定、同步或解绑"],
            ["页面实测身份", "企业主账号负责人"],
        ],
        [1800, 7560],
        PALE_BLUE,
    )

    doc.add_heading("5.1 功能概述", level=1)
    add_body(
        doc,
        "账号与访问管理用于建立Harson-Base登录身份、内部AIGC账号关系与CL-AIGC工作区会话之间的对应关系。平台将Harson-Base用户作为统一登录主体，再按管理员、企业主账号负责人、已绑定普通用户和未绑定普通用户四类访问状态分配页面入口与数据范围。",
    )
    add_body(
        doc,
        "系统中的“企业主账号”和“内部AIGC子账号”均为平台内部管理对象；主账号负责人和普通用户仍使用各自的Harson-Base邮箱及密码登录。进入创作模块时，后端根据账号映射选择主账号外部会话或子账号独立外部会话，页面不单独回显登录Token。",
    )

    add_table(
        doc,
        ["身份/状态", "可见入口与数据范围", "核实结论"],
        [
            ["管理员", "账号管理全量功能；企业主账号、外部账号、子账号、负责人和用户映射管理", "代码已核实；管理员页面未在本次会话实测"],
            ["企业主账号负责人", "企业仪表盘、本企业子账号展示、CL-AIGC工作区", "代码与浏览器均已核实"],
            ["已绑定普通用户", "个人绑定信息、本人创作记录、所映射子账号的工作区", "代码已核实；普通用户页面待实测"],
            ["未绑定普通用户", "账号页显示未开通提示；不能建立CL-AIGC会话", "代码已核实；页面待实测"],
            ["未登录访客", "登录/注册页；访问工作区时重定向登录页", "浏览器已实测"],
        ],
        [1650, 4780, 2930],
    )

    doc.add_heading("5.2 用户登录与访问", level=1)
    doc.add_heading("5.2.1 登录入口与字段", level=2)
    add_body(
        doc,
        "入口：用户可在Harson-Base首页点击右上角用户图标进入“登录 Harson-Base”页面，也可直接访问/login。未登录用户直接访问CL-AIGC工作区时，系统会自动重定向至该登录页。",
        bold_prefix="入口：",
    )
    add_table(
        doc,
        ["区域", "字段/控件及用途"],
        [
            ["登录", "邮箱（email，必填）、密码（password，必填）、“登录”按钮"],
            ["注册", "姓名、邮箱、密码（前端最少8位）、性别、年龄组及“创建账户”按钮"],
            ["辅助入口", "“返回首页”链接"],
        ],
        [1800, 7560],
    )
    # Keep the login figure on the same page as its field table.
    add_picture(doc, LOGIN_IMAGE, 5.4)
    add_caption(doc, "图5-1  Harson-Base登录/注册页面（本地页面实测，未输入任何账号凭据）")

    doc.add_heading("5.2.2 登录操作与处理结果", level=2)
    add_step(doc, "一", "输入登录信息", "在登录区填写Harson-Base邮箱和密码。页面仅将两项内容提交至站内登录接口，不在界面显示密码明文。")
    add_step(doc, "二", "校验用户", "后端按邮箱查找启用状态的用户，并使用密码哈希进行比对。用户不存在、已停用或密码不匹配时，返回“邮箱或密码无效”，不建立登录态。")
    add_step(doc, "三", "建立登录态", "校验成功后更新最近登录时间，并签发包含用户ID、邮箱和角色的登录令牌。令牌写入harson_token Cookie；Cookie设置为HttpOnly、SameSite=Strict，生产环境启用Secure，默认有效期为2小时。")
    add_step(doc, "四", "返回首页", "前端显示登录成功结果并跳转至平台首页。首页根据当前身份显示用户头像；具备权限时显示“仪表盘”入口，所有登录用户均可继续访问受保护的工作区路由。")

    add_note(
        doc,
        "安全说明",
        "本章不记录真实邮箱、密码、Cookie值或Token值。浏览器截图未包含登录凭据；工作区截图已避开账户余额区域。",
    )

    doc.add_heading("5.2.3 页面访问控制", level=2)
    add_table(
        doc,
        ["访问目标", "校验方式", "允许结果", "拒绝结果"],
        [
            ["CL-AIGC工作区", "校验harson_token及用户ID", "返回工作区页面", "无Cookie或令牌失效时重定向/login"],
            ["账号管理", "页面依次尝试管理员、主账号负责人、普通用户数据接口", "按身份渲染对应视图", "未登录时显示“请先登录”及“前往登录”"],
            ["企业仪表盘", "管理员角色或有效主账号负责人绑定", "管理员查看全量；负责人仅查看所属企业", "无权限时不返回仪表盘内容；数据接口返回禁止访问"],
            ["管理接口", "先校验登录，再校验admin角色", "管理员可执行", "非管理员返回无管理员权限"],
        ],
        [1700, 2600, 2550, 2510],
    )

    doc.add_heading("5.2.4 工作区访问流程", level=2)
    add_body(
        doc,
        "工作区入口：平台首页“进入 AIGC 平台”、企业仪表盘“进入 CL-AIGC”以及账号管理页“进入 CL-AIGC”均指向/aigc-workspace。页面先展示CL-AIGC自有首页、模块导航、知识库入口、积分入口和刷新按钮。",
        bold_prefix="工作区入口：",
    )
    add_step(doc, "一", "进入工作区", "服务端先校验Harson-Base登录Cookie。校验成功后返回本地CL-AIGC工作区；校验失败则跳转登录页。")
    add_step(doc, "二", "选择功能模块", "用户点击“图片创作”“高清放大”“视频创作”“AI图案设计”“提示词生成”“AI服装”或“AI电商”等模块。仅停留在本地首页时，不需要推测外部模块页面按钮。")
    add_step(doc, "三", "解析账号关系", "企业主账号负责人通过负责人绑定定位企业主账号；普通用户通过Harson-Base用户映射定位内部AIGC子账号及所属企业。若同一用户同时存在负责人绑定和子账号映射，系统返回关系冲突错误。")
    add_step(doc, "四", "建立外部会话", "负责人继续使用所属企业主账号的外部会话；普通用户使用内部子账号对应的外部子账号独立会话。系统优先验证缓存Token，失效或不存在时才使用后端保存的加密凭据重新登录。")
    add_step(doc, "五", "加载模块", "后端按外部菜单名称解析真实路由，生成同源白名单范围内的嵌入地址，并以iframe加载模块。Token仅包含在后端生成的iframe地址中，不作为独立响应字段展示。")
    # Fit the workspace evidence with the access-flow steps and pending note.
    add_picture(doc, WORKSPACE_IMAGE, 4.0)
    add_caption(doc, "图5-2  企业主账号负责人登录后的CL-AIGC本地工作区（脱敏截图，余额区域未纳入画面）")
    add_note(
        doc,
        "待核实",
        "本次为避免触发外部账号重新登录、任务同步或生产会话变更，未点击具体创作模块进入外部iframe；外部模块内页面和按钮行为应在专用测试账号环境中另行核实。",
        "pending",
    )

    doc.add_heading("5.3 企业主账号管理", level=1)
    doc.add_heading("5.3.1 管理入口与内部主账号", level=2)
    add_body(
        doc,
        "入口：管理员登录后进入“账号管理”。页面完整管理视图按顺序提供内部企业主账号、外部主账号绑定、内部子账号、外部子账号、Token配额、企业负责人及Harson-Base用户映射功能。非管理员不能调用上述新增或绑定接口。",
        bold_prefix="入口：",
    )
    add_table(
        doc,
        ["字段", "填写要求与作用"],
        [
            ["企业名称", "必填，用于标识所属企业"],
            ["平台名称", "必填，例如外部AIGC平台名称"],
            ["内部主账号管理邮箱", "必填、邮箱格式；作为内部企业主账号登录标识，系统要求不重复"],
            ["内部主账号独立密码", "必填；仅保存密码哈希，不在管理列表回显"],
            ["初始总点数", "界面隐藏并固定为0；完成外部主账号绑定后再同步真实点数"],
        ],
        [2500, 6860],
    )
    add_step(doc, "一", "填写主账号资料", "管理员在“创建 AIGC 企业主账号”区域填写企业名称、平台名称、内部管理邮箱和独立密码。")
    add_step(doc, "二", "提交创建", "系统检查必填项和登录邮箱唯一性，对密码执行哈希处理，生成内部主账号ID，并将状态设为active。")
    add_step(doc, "三", "查看结果", "创建成功后返回“AIGC 企业主账号创建成功”，主账号进入管理下拉列表；此时总点数仍为0，尚不能代表外部真实余额。")

    doc.add_heading("5.3.2 外部主账号绑定与点数同步", level=2)
    add_table(
        doc,
        ["字段", "填写要求与处理"],
        [
            ["内部AIGC企业主账号", "从已创建的启用主账号中选择"],
            ["外部账号", "必填，用于向CL-AIGC外部服务验证登录"],
            ["外部登录密码", "必填；前端提交后清空输入框，后端验证成功后加密保存，不回显"],
            ["点数字段", "从balance、companyBalance、mpoint、companyMpoint中选择余额来源"],
        ],
        [2500, 6860],
    )
    add_step(doc, "一", "选择内部主账号", "在“绑定外部账号并同步总点数”区域选择目标企业主账号。")
    add_step(doc, "二", "验证外部凭据", "输入外部账号和密码后点击“验证账号并完成绑定”。后端先调用外部登录接口；验证失败时不建立绑定。")
    add_step(doc, "三", "保存并同步", "验证成功后加密保存外部凭据，记录外部企业及成员信息，按所选点数字段同步主账号总点数，并保存主账号级登录缓存。")
    add_step(doc, "四", "后续维护", "绑定列表提供“重新同步点数”“同步创作记录”和“解绑”操作。解绑会清除外部绑定、登录缓存、真实任务快照，并将主账号及其子账号点数归零；内部主账号、内部子账号和Harson-Base映射仍保留。")
    add_note(
        doc,
        "待核实",
        "管理员完整管理界面、外部凭据验证及点数同步行为已由代码核实，本次未使用管理员账号或外部生产账号实际提交。",
        "pending",
    )

    doc.add_heading("5.3.3 企业主账号负责人绑定", level=2)
    add_table(
        doc,
        ["字段", "规则"],
        [
            ["企业主账号", "只能选择尚未绑定负责人的主账号"],
            ["Harson-Base负责人", "只能选择非管理员、未绑定子账号、且不是其他企业负责人的启用用户"],
        ],
        [2500, 6860],
    )
    add_step(doc, "一", "选择账号", "管理员分别选择企业主账号和Harson-Base负责人。")
    add_step(doc, "二", "建立负责人关系", "系统校验一主账号一负责人、一负责人一主账号，并阻止管理员账号或已绑定子账号的用户成为负责人。")
    add_step(doc, "三", "获得访问范围", "绑定成功后，该负责人可看到企业仪表盘入口；仪表盘查询被限制为所属主账号，不能查看或同步其他企业数据。")
    add_step(doc, "四", "负责人账号管理页", "负责人进入“账号管理”时，仅显示本企业子账号。当前页面明确标注“演示数据”，其中“修改额度”不会写入系统数据；真实额度修改仍属于管理员功能。")
    add_step(doc, "五", "解除负责人绑定", "管理员在负责人总览点击“解除负责人绑定”并确认。关系状态改为disabled后，该用户失去企业负责人仪表盘权限。")
    add_note(doc, "页面实测", "本次使用企业主账号负责人身份验证了企业子账号管理页、企业仪表盘和CL-AIGC工作区入口；未执行任何修改、同步或解绑操作。")

    doc.add_heading("5.4 内部子账号与用户关联", level=1)
    doc.add_heading("5.4.1 创建内部AIGC子账号", level=2)
    add_table(
        doc,
        ["字段", "填写要求与校验"],
        [
            ["企业主账号", "必选，且主账号状态必须为active"],
            ["子账号名称", "必填，用于组织内识别"],
            ["AIGC子账号登录邮箱", "必填、邮箱格式，内部子账号之间不可重复"],
            ["AIGC子账号独立密码", "必填；只保存密码哈希"],
            ["Token配额", "必填、不得为负；与同一主账号其他子账号配额合计不得超过主账号总点数"],
            ["剩余预警阈值", "可填写1—100的百分比，默认10；剩余比例小于或等于阈值时进入预警状态"],
        ],
        [2500, 6860],
    )
    add_step(doc, "一", "填写子账号信息", "管理员选择所属企业主账号，填写名称、登录邮箱、独立密码、Token配额和剩余预警阈值。")
    add_step(doc, "二", "校验企业容量", "系统确认主账号存在且启用，并计算同一企业已分配配额；本次配额不能使分配总额超过主账号总点数。")
    add_step(doc, "三", "保存子账号", "系统对密码哈希后创建active状态的内部子账号，并返回配额、已用量、剩余量、使用率和预警状态。")

    doc.add_heading("5.4.2 绑定外部子账号", level=2)
    add_body(
        doc,
        "管理员页面通过附加的“绑定外部子账号”区域，将内部AIGC子账号与外部平台成员账号关联。绑定成功后，映射到该内部子账号的普通Harson-Base用户使用该外部子账号的独立会话；企业主账号负责人仍使用主账号会话。",
    )
    add_table(
        doc,
        ["字段", "填写要求与结果"],
        [
            ["AIGC子账号", "选择一个启用的内部子账号；已绑定项可重新验证更新"],
            ["外部子账号登录账号", "必填；不能与所属外部主账号相同，也不能已绑定给其他内部子账号"],
            ["外部子账号登录密码", "必填；验证成功后加密保存"],
            ["点数字段", "选择balance、companyBalance、mpoint或companyMpoint"],
        ],
        [2500, 6860],
    )
    add_step(doc, "一", "选择并验证", "管理员选择内部子账号，填写外部子账号凭据，点击“验证并绑定外部子账号”。")
    add_step(doc, "二", "核对企业归属", "后端验证外部登录成功后，确认该成员属于当前主账号对应的外部企业，并记录外部成员ID和点数。")
    add_step(doc, "三", "建立独立会话", "普通用户访问工作区时，系统以内部子账号ID读取子账号级缓存；缓存不可用时才解密凭据重新登录。")
    add_step(doc, "四", "同步或解绑", "管理员可在绑定列表执行“同步”或“解绑”。解绑外部子账号时，内部子账号和Harson-Base用户映射保留，但普通用户将无法建立外部子账号会话，直至重新绑定。")
    add_note(
        doc,
        "待核实",
        "外部子账号绑定、同步及解绑已由代码核实，本次未使用管理员账号提交外部凭据，相关成功提示和外部返回数据待专用测试环境实测。",
        "pending",
    )

    doc.add_heading("5.4.3 Harson-Base用户映射", level=2)
    add_body(
        doc,
        "入口：管理员在账号管理页“建立 Harson-Base ↔ AIGC 子账号映射”区域操作。",
        bold_prefix="入口：",
    )
    add_table(
        doc,
        ["字段", "规则"],
        [
            ["Harson-Base用户", "仅显示非管理员、尚无有效子账号映射、且不是企业主账号负责人的用户"],
            ["AIGC子账号", "选择一个已创建且启用的内部AIGC子账号"],
        ],
        [2500, 6860],
    )
    add_step(doc, "一", "选择关联双方", "管理员从下拉列表选择Harson-Base用户和内部AIGC子账号。")
    add_step(doc, "二", "创建映射", "系统确认Harson-Base用户存在、内部子账号启用，并确认该用户当前没有其他有效子账号映射。")
    add_step(doc, "三", "保存关系", "系统记录用户ID、登录邮箱、内部子账号ID、所属主账号ID和active状态，并返回“一对一映射创建成功”。这里的“一对一”表示一个Harson-Base用户只能绑定一个子账号；当前代码允许同一内部子账号被多个Harson-Base用户使用。")
    add_step(doc, "四", "用户访问", "普通用户登录后只能查看自己的绑定信息和与该子账号身份匹配的创作记录；进入工作区时使用该内部子账号绑定的外部子账号会话。")
    add_step(doc, "五", "解除映射", "管理员在映射总览点击“解除映射”并确认。映射状态改为disabled，系统同时尝试清理该用户的AIGC会话；内部主账号、内部子账号和历史创作记录不删除。")

    doc.add_heading("5.4.4 未绑定与异常关系处理", level=2)
    add_table(
        doc,
        ["状态", "页面或系统结果"],
        [
            ["用户未绑定任何AIGC账号", "账号管理页显示“尚未开通AIGC服务”；不能进入实际创作会话"],
            ["内部子账号不存在或停用", "拒绝建立工作区会话"],
            ["所属主账号不存在或停用", "拒绝建立工作区会话"],
            ["普通用户子账号未绑定外部子账号", "提示尚未绑定外部子账号，不生成iframe会话"],
            ["同一用户同时存在负责人和子账号映射", "返回关系冲突错误，要求管理员修复映射"],
            ["主账号、子账号与映射记录的企业ID不一致", "返回关系不一致错误，拒绝继续"],
        ],
        [3200, 6160],
    )

    doc.add_heading("5.5 退出登录", level=1)
    add_body(
        doc,
        "入口：首页登录头像、账号管理页“退出登录”按钮和企业仪表盘右上角“退出”按钮均可发起退出。",
        bold_prefix="入口：",
    )
    add_step(doc, "一", "发起退出", "前端向站内退出接口发送POST请求，并携带当前Harson-Base登录Cookie。")
    add_step(doc, "二", "标记AIGC会话", "系统将当前用户的AIGC活跃会话标记为已退出。企业负责人按企业主账号共享会话处理；普通用户按内部子账号共享会话处理。")
    add_step(doc, "三", "判断是否保留共享Token", "如果同一企业主账号或同一内部子账号仍有其他活跃Harson-Base用户，系统保留共享Token；当前用户退出不影响其他用户。")
    add_step(doc, "四", "注销外部会话", "如果已无其他活跃用户且存在缓存Token，系统调用外部注销接口。成功后清除相应Token缓存；若本地缓存损坏或不存在，则仅清理本地会话记录。")
    add_step(doc, "五", "清除本地登录态", "无论外部AIGC注销是否成功，系统最终都会清除harson_token Cookie，并返回“Harson-Base 已退出登录”。外部注销失败仅作为警告，不阻止本地退出。")
    add_step(doc, "六", "页面跳转", "首页、账号管理页和仪表盘退出控件均在请求结束后返回平台首页。退出后再次访问/aigc-workspace会被重定向至/login。")
    add_note(doc, "页面实测", "本次已在企业仪表盘点击“退出”，页面返回平台首页；随后直接访问CL-AIGC工作区，实际重定向至登录页。未展示任何Cookie或外部Token。")

    doc.add_heading("5.6 核实结论与待核实项", level=1)
    add_table(
        doc,
        ["核实对象", "结论", "状态"],
        [
            ["登录页字段、登录后首页入口", "与代码一致", "已实测"],
            ["未登录访问工作区", "重定向/login", "已实测"],
            ["企业主账号负责人账号页", "仅展示本企业子账号，并标注额度操作为演示", "已实测"],
            ["企业负责人仪表盘和工作区", "可访问，范围限制为所属企业", "已实测"],
            ["退出登录及退出后拦截", "返回首页；再次访问工作区重定向登录页", "已实测"],
            ["管理员创建/绑定/同步/解绑", "字段、校验和后端结果已按代码整理", "待管理员测试账号实测"],
            ["已绑定普通用户个人账号页", "字段与数据隔离逻辑已按代码整理", "待普通测试账号实测"],
            ["外部创作模块iframe内功能", "会话与路由生成逻辑已按代码整理", "待专用测试环境联调"],
        ],
        [2850, 4460, 2050],
    )
    add_note(
        doc,
        "说明",
        "本章结论以2026年9月24日工作区代码和本地页面为准。未实测项目均已明确标注，不根据页面文案推测未触发按钮的实际结果。",
    )

    doc.core_properties.title = "CL-AIGC智能创作与管理平台V1.0设计说明书——第5章 账号与访问管理"
    doc.core_properties.subject = "软件著作权申请设计说明书"
    doc.core_properties.author = "HARSON"
    doc.core_properties.keywords = "CL-AIGC, Harson-Base, 账号管理, 访问控制"
    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    print(build_document())
