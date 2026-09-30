from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "references" / "CL-AIGC智能创作与管理平台V1.0设计说明书-第5章-申报版.docx"
LOGIN_IMAGE = ROOT / "references" / "chapter5-screenshots" / "login.png"

FONT = "Arial Unicode MS"
BLACK = "000000"
NAVY = "203864"
LIGHT_BLUE = "D9EAF7"
ALT_BLUE = "F3F7FB"
BORDER = "D9D9D9"
GRAY = "666666"
CONTENT_DXA = 9360


def set_run(run, size=10.5, bold=False, color=BLACK):
    run.font.name = FONT
    rpr = run._element.get_or_add_rPr()
    for name in ("ascii", "hAnsi", "eastAsia"):
        rpr.rFonts.set(qn(f"w:{name}"), FONT)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=85, start=110, bottom=85, end=110):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "6")
        node.set(qn("w:color"), BORDER)


def set_table_geometry(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            tc_w.set(qn("w:w"), str(widths[idx]))
            tc_w.set(qn("w:type"), "dxa")
    set_borders(table)


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths)
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_shading(cell, LIGHT_BLUE)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        set_run(p.add_run(header), 9.5, True)
    for row_index, values in enumerate(rows):
        row = table.add_row()
        if row_index % 2:
            for cell in row.cells:
                set_cell_shading(cell, ALT_BLUE)
        for idx, value in enumerate(values):
            p = row.cells[idx].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            set_run(p.add_run(str(value)), 9.5, idx == 0 and len(headers) == 2)
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_body(doc, text, lead=None):
    p = doc.add_paragraph(style="Normal")
    if lead and text.startswith(lead):
        set_run(p.add_run(lead), bold=True)
        set_run(p.add_run(text[len(lead):]))
    else:
        set_run(p.add_run(text))
    return p


def add_step(doc, number, title, detail):
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.left_indent = Inches(0.16)
    p.paragraph_format.first_line_indent = Inches(-0.16)
    p.paragraph_format.keep_together = True
    set_run(p.add_run(f"步骤{number}  {title}  "), bold=True)
    set_run(p.add_run(detail))
    return p


def add_heading(doc, text, level):
    p = doc.add_paragraph(style=f"Heading {level}")
    set_run(p.add_run(text), {1: 17, 2: 14, 3: 12}[level], True)
    return p


def configure(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.78)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.88)
    section.right_margin = Inches(0.88)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal.paragraph_format.line_spacing = 1.18
    normal.paragraph_format.space_after = Pt(5)

    for style_name, size, before, after in (
        ("Heading 1", 17, 12, 8),
        ("Heading 2", 14, 10, 6),
        ("Heading 3", 12, 8, 4),
    ):
        style = doc.styles[style_name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLACK)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    title_style = doc.styles["Title"]
    title_style.font.name = FONT
    title_style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    title_style.font.size = Pt(22)
    title_style.font.bold = True
    title_style.font.color.rgb = RGBColor.from_string(BLACK)
    title_ppr = title_style._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_run(header.add_run("CL-AIGC智能创作与管理平台V1.0设计说明书"), 8.5, False, BLACK)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run(footer.add_run("第5章 账号与访问管理  "), 8.5, False, BLACK)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    footer._p.append(fld)


def build():
    doc = Document()
    configure(doc)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.paragraph_format.space_after = Pt(10)
    set_run(title.add_run("第5章 账号与访问管理"), 22, True)

    add_heading(doc, "5.1 功能概述", 2)
    add_body(doc, "账号与访问管理为CL-AIGC智能创作与管理平台提供统一身份认证、企业账号组织、内部子账号分配和工作区访问控制。平台使用Harson-Base作为统一身份认证入口，用户完成认证后，系统根据用户角色和账号映射关系确定可访问的页面、企业数据及AIGC工作区会话。")
    add_body(doc, "系统将账号分为管理员、企业主账号负责人、已关联内部子账号的普通用户和未关联账号的普通用户。管理员维护企业主账号、内部子账号及用户关联；企业主账号负责人查看所属企业的数据和子账号；普通用户使用本人关联的内部子账号访问工作区。")
    add_table(doc, ["用户类型", "主要权限", "数据范围"], [
        ["管理员", "维护企业主账号、外部账号绑定、内部子账号、负责人及用户关联", "全部企业账号管理数据"],
        ["企业主账号负责人", "访问企业仪表盘、查看本企业子账号、进入CL-AIGC工作区", "所属企业"],
        ["已关联普通用户", "查看本人关联信息和创作记录、进入CL-AIGC工作区", "本人及关联子账号"],
        ["未关联普通用户", "登录平台并查看账号开通提示", "本人基础信息"],
    ], [1900, 4660, 2800])

    add_heading(doc, "5.2 用户登录与访问", 2)
    add_heading(doc, "5.2.1 登录入口和字段", 3)
    add_body(doc, "用户可从平台首页进入登录页面，也可在访问受保护页面时由系统自动跳转至登录页面。Harson-Base是本平台的统一身份认证入口，登录成功后用户可继续访问CL-AIGC平台功能。")
    add_table(doc, ["区域", "字段或控件", "说明"], [
        ["登录", "邮箱、密码、登录按钮", "邮箱和密码均为必填项"],
        ["注册", "姓名、邮箱、密码、性别、年龄组、创建账户按钮", "密码长度不得少于8位"],
        ["页面导航", "返回首页", "返回平台公共首页"],
    ], [1500, 3860, 4000])

    if LOGIN_IMAGE.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(str(LOGIN_IMAGE), width=Inches(4.6))
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.space_after = Pt(6)
        set_run(cap.add_run("图5-1  Harson-Base统一登录和注册页面"), 9, False, GRAY)

    add_heading(doc, "5.2.2 登录处理流程", 3)
    add_step(doc, "一", "输入登录信息", "用户填写Harson-Base邮箱和密码并提交登录请求。密码字段采用受保护输入方式，页面不显示密码明文。")
    add_step(doc, "二", "校验用户", "系统按邮箱查询启用状态的用户，并校验密码。用户不存在、账号停用或密码不匹配时，系统提示邮箱或密码无效。")
    add_step(doc, "三", "建立登录状态", "校验成功后，系统记录最近登录时间，建立包含用户标识、邮箱和角色的登录状态，并通过受保护的Cookie维持会话。")
    add_step(doc, "四", "进入平台", "页面显示登录成功信息并返回平台首页。系统根据用户身份显示对应入口，用户可继续访问账号管理、企业仪表盘或CL-AIGC工作区。")

    add_heading(doc, "5.2.3 页面访问控制", 3)
    add_table(doc, ["访问目标", "访问条件", "处理结果"], [
        ["CL-AIGC工作区", "用户已登录", "返回工作区页面；未登录时跳转至登录页面"],
        ["账号管理", "用户已登录", "按管理员、企业负责人或普通用户身份显示对应视图"],
        ["企业仪表盘", "管理员或有效的企业主账号负责人", "管理员可查看全部企业；负责人仅查看所属企业"],
        ["管理接口", "用户已登录且角色为管理员", "校验通过后执行；非管理员请求被拒绝"],
    ], [2100, 3400, 3860])

    add_heading(doc, "5.2.4 工作区访问流程", 3)
    add_step(doc, "一", "进入工作区", "用户从平台首页、企业仪表盘或账号管理页进入CL-AIGC工作区。服务端首先校验当前登录状态。")
    add_step(doc, "二", "选择功能模块", "工作区显示图片创作、视频创作、AI图案设计、提示词生成、AI服装和AI电商等功能入口。")
    add_step(doc, "三", "解析账号关系", "企业负责人通过负责人关系定位企业主账号；普通用户通过Harson-Base用户关联定位内部子账号及所属企业。")
    add_step(doc, "四", "建立AIGC会话", "企业负责人使用所属企业主账号会话；普通用户使用其内部子账号对应的外部子账号独立会话。")
    add_step(doc, "五", "加载功能页面", "系统按功能名称解析可用路由，在平台允许的地址范围内加载相应模块。用户无需再次输入外部平台账号。")

    add_heading(doc, "5.3 企业主账号管理", 2)
    add_heading(doc, "5.3.1 企业主账号", 3)
    add_body(doc, "管理员进入账号管理页面后，可以创建和维护企业主账号。企业主账号用于表示一个企业在CL-AIGC平台中的组织主体，并作为内部子账号、企业负责人和外部平台账号的归属对象。")
    add_table(doc, ["字段", "填写要求及作用"], [
        ["企业名称", "必填，用于标识账号所属企业"],
        ["平台名称", "必填，用于记录所关联的AIGC平台"],
        ["内部主账号管理邮箱", "必填且符合邮箱格式，在企业主账号中不得重复"],
        ["内部主账号独立密码", "必填，系统进行不可逆处理后保存，不在管理列表中回显"],
        ["初始总点数", "创建时为0，完成外部主账号绑定后同步实际点数"],
    ], [2900, 6460])
    add_step(doc, "一", "填写资料", "管理员填写企业名称、平台名称、内部管理邮箱和独立密码。")
    add_step(doc, "二", "提交创建", "系统校验必填项和邮箱唯一性，创建企业主账号并将状态设置为启用。")
    add_step(doc, "三", "查看结果", "创建成功后，企业主账号进入账号管理列表，可继续配置外部主账号、内部子账号和企业负责人。")

    add_heading(doc, "5.3.2 外部主账号绑定", 3)
    add_table(doc, ["字段", "填写要求及作用"], [
        ["企业主账号", "从已创建且处于启用状态的企业主账号中选择"],
        ["外部账号", "必填，用于验证企业在外部AIGC平台中的登录身份"],
        ["外部登录密码", "必填，验证成功后加密保存，页面不回显"],
        ["点数字段", "选择外部平台中用于表示企业点数余额的字段"],
    ], [2900, 6460])
    add_step(doc, "一", "选择企业主账号", "管理员选择需要关联外部平台账号的企业主账号。")
    add_step(doc, "二", "验证外部账号", "管理员填写外部账号和密码。系统调用外部登录服务验证账号，验证失败时不保存绑定关系。")
    add_step(doc, "三", "保存绑定", "验证成功后，系统保存加密凭据、外部企业标识和成员信息，并同步企业主账号点数。")
    add_step(doc, "四", "维护绑定", "管理员可执行点数同步、创作记录同步和解绑。解绑后保留内部主账号、内部子账号及用户关联。")

    add_heading(doc, "5.3.3 企业主账号负责人", 3)
    add_body(doc, "管理员可将一个Harson-Base用户设置为企业主账号负责人。企业主账号只能配置一名有效负责人，同一用户不能同时担任多个企业的负责人，也不能同时作为普通用户关联内部子账号。管理员账号不能设置为企业负责人。")
    add_step(doc, "一", "选择账号", "管理员选择企业主账号和符合条件的Harson-Base用户。")
    add_step(doc, "二", "建立负责人关系", "系统校验主账号和用户状态，确认不存在重复负责人或子账号关联后保存关系。")
    add_step(doc, "三", "分配访问范围", "负责人登录后可进入企业仪表盘和CL-AIGC工作区，查询范围限定为所属企业。")
    add_step(doc, "四", "解除负责人关系", "管理员确认解除后，负责人关系被停用，该用户不再具有企业负责人访问权限。")

    add_heading(doc, "5.4 内部子账号与用户关联", 2)
    add_heading(doc, "5.4.1 创建内部子账号", 3)
    add_body(doc, "内部子账号用于划分企业内部的AIGC使用身份和点数额度。每个内部子账号必须归属于一个启用状态的企业主账号。")
    add_table(doc, ["字段", "填写要求及校验"], [
        ["企业主账号", "必选，且状态必须为启用"],
        ["子账号名称", "必填，用于企业内部识别"],
        ["子账号登录邮箱", "必填且符合邮箱格式，内部子账号之间不得重复"],
        ["子账号独立密码", "必填，系统处理后保存，不在管理列表中回显"],
        ["Token配额", "不得为负；同一企业的已分配配额合计不得超过企业总点数"],
        ["剩余预警阈值", "取值范围为1至100，默认值为10，用于判断剩余比例预警"],
    ], [2900, 6460])
    add_step(doc, "一", "填写子账号资料", "管理员选择所属企业主账号，填写名称、登录邮箱、独立密码、Token配额和预警阈值。")
    add_step(doc, "二", "校验企业容量", "系统确认企业主账号存在并处于启用状态，计算企业已分配配额，防止分配总额超过企业总点数。")
    add_step(doc, "三", "保存子账号", "校验通过后创建启用状态的内部子账号，并显示配额、已用量、剩余量、使用率和预警状态。")

    add_heading(doc, "5.4.2 外部子账号绑定", 3)
    add_body(doc, "管理员可以将内部子账号与外部AIGC平台成员账号关联。完成关联后，普通用户通过内部子账号使用对应的外部子账号独立会话。")
    add_table(doc, ["字段", "填写要求及处理结果"], [
        ["内部子账号", "选择一个启用状态的内部子账号"],
        ["外部子账号", "不得与所属企业的外部主账号相同，也不得已关联其他内部子账号"],
        ["外部子账号密码", "验证成功后加密保存，页面不回显"],
        ["点数字段", "选择用于表示外部子账号余额的字段"],
    ], [2900, 6460])
    add_step(doc, "一", "验证账号", "管理员选择内部子账号并填写外部子账号凭据，系统验证外部登录信息。")
    add_step(doc, "二", "核对企业归属", "系统确认外部成员属于当前企业主账号对应的外部企业，并记录外部成员标识和点数。")
    add_step(doc, "三", "建立独立会话", "普通用户进入工作区时，系统按内部子账号建立或恢复子账号级会话。")
    add_step(doc, "四", "同步或解绑", "管理员可同步外部子账号信息或解除绑定。解绑不删除内部子账号和Harson-Base用户关联。")

    add_heading(doc, "5.4.3 Harson-Base用户关联", 3)
    add_body(doc, "管理员在账号管理页面选择Harson-Base用户和内部子账号，建立登录用户与AIGC使用身份之间的关联。可选用户必须为非管理员、没有其他有效子账号关联且不是企业主账号负责人的启用用户。")
    add_step(doc, "一", "选择关联双方", "管理员选择Harson-Base用户和一个启用状态的内部子账号。")
    add_step(doc, "二", "校验关联条件", "系统确认用户和内部子账号存在，检查该用户没有其他有效子账号关联。")
    add_step(doc, "三", "保存关联", "系统记录用户标识、登录邮箱、内部子账号标识、所属企业主账号标识和启用状态。一个Harson-Base用户只能关联一个内部子账号；同一内部子账号可以供多个Harson-Base用户使用。")
    add_step(doc, "四", "用户访问", "普通用户登录后可查看本人关联信息和与该子账号身份匹配的创作记录，进入工作区时使用该子账号对应的外部会话。")
    add_step(doc, "五", "解除关联", "管理员确认解除后，关联状态被停用，系统清理该用户的AIGC会话。内部账号和历史创作记录继续保留。")

    add_heading(doc, "5.5 异常处理", 2)
    add_table(doc, ["异常状态", "系统处理"], [
        ["用户未关联AIGC账号", "账号管理页显示尚未开通AIGC服务，不建立实际创作会话"],
        ["内部子账号不存在或停用", "拒绝建立工作区会话"],
        ["所属企业主账号不存在或停用", "拒绝建立工作区会话"],
        ["外部子账号未绑定", "提示尚未绑定外部子账号，不加载外部功能页面"],
        ["同一用户同时存在负责人关系和子账号关联", "返回账号关系冲突，要求管理员修正关联"],
        ["企业主账号、内部子账号与关联记录的企业标识不一致", "返回关系不一致，拒绝继续访问"],
        ["登录状态失效", "清除无效登录状态并跳转至登录页面"],
    ], [3550, 5810])

    add_heading(doc, "5.6 退出登录", 2)
    add_body(doc, "用户可以通过平台首页、账号管理页或企业仪表盘的退出控件结束当前登录状态。退出处理同时清理Harson-Base登录状态和当前用户对应的AIGC活跃会话。")
    add_step(doc, "一", "发起退出", "前端向平台退出接口发送请求，并携带当前登录状态。")
    add_step(doc, "二", "标记AIGC会话", "系统将当前用户对应的AIGC活跃会话标记为已退出。企业负责人按企业主账号处理，普通用户按内部子账号处理。")
    add_step(doc, "三", "检查共享会话", "如果同一企业主账号或内部子账号仍有其他活跃用户，系统保留共享会话，避免影响其他用户。")
    add_step(doc, "四", "注销外部会话", "不存在其他活跃用户时，系统注销相应的外部会话并清理会话缓存。外部注销异常不阻止本地退出。")
    add_step(doc, "五", "清除本地登录状态", "系统清除当前用户的Harson-Base登录Cookie，并向前端返回退出结果。")
    add_step(doc, "六", "返回首页", "页面返回平台首页。用户再次访问CL-AIGC工作区时，系统将其跳转至登录页面。")

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
